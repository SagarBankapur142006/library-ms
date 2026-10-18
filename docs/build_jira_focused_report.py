from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "docs" / "jira_focused_report_assets"
SNAPSHOT = ASSET_DIR / "jira_snapshot.json"
OUT = ROOT / "docs" / "Jira_Focused_Library_Management_System_Report.docx"


PROJECT_DESCRIPTIONS = {
    "SCRUM-1": "Initial Jira starter work item used to add the team's work and attach early implementation evidence.",
    "SCRUM-2": "Configured connection between the coding workflow and Jira so project activity could be tracked from the board.",
    "SCRUM-3": "Set up Atlassian Rovo/MCP integration support for local coding-tool interaction.",
    "SCRUM-4": "Subtask for delegating a Jira work item to a coding agent.",
    "SCRUM-5": "Subtask for implementing the delegated Jira work item from the development environment.",
    "SCRUM-6": "Database design task covering schema modelling, normalized entities, and ER diagram evidence.",
    "SCRUM-7": "Login page UI task for the authentication entry screen and visual interface proof.",
    "SCRUM-8": "Authentication implementation task covering the React authentication context.",
    "SCRUM-9": "Book management design task for catalog browsing and inventory-management UI.",
    "SCRUM-10": "Requirement research task covering WBS, SRS, use-case analysis, and early planning documents.",
    "SCRUM-11": "Project structure task covering the application codebase and folder/document organization.",
    "SCRUM-12": "Project Planning epic containing planning document evidence and handwritten planning photographs.",
    "SCRUM-13": "User Authentication epic for login, role-aware sessions, and access-control work.",
    "SCRUM-14": "Book Management epic for catalog CRUD, search, filtering, and book availability handling.",
    "SCRUM-15": "Member Management epic covering student/member loan and reservation visibility.",
    "SCRUM-16": "Issue & Return epic covering checkout, return, fine, reservation, and circulation logic.",
    "SCRUM-17": "Reports & Dashboard epic covering dashboard metrics and settings evidence.",
    "SCRUM-18": "Requirement Analysis epic covering formal SRS and use-case diagram evidence.",
    "SCRUM-19": "System Design epic covering ER modelling and database schema documentation.",
}


ATTACHMENT_DESCRIPTIONS = {
    "add-your-teams-work.png": "Jira starter-screen evidence showing how the team's work item was initiated.",
    "implementation_report.md": "Short implementation note uploaded as supporting work evidence.",
    "main_dashboard_1785261292483.png": "Dashboard UI screenshot showing system metrics and implementation progress.",
    "Project_Documentation.html": "Compiled project documentation page attached as implementation evidence.",
    "connect-agents-to-jira.png": "Screenshot showing Jira coding-agent connection instructions/evidence.",
    "connect-atlassian-mcp.png": "Screenshot showing Atlassian MCP/Rovo connection instructions/evidence.",
    "Database_Design.md": "Database design markdown describing tables, constraints, relationships, and normalization.",
    "er_diagram.png": "StarUML-style ER diagram image showing entities and relationships for the LMS database.",
    "System_Design_Database_Schema.docx": "System design document containing database schema and design explanation.",
    "Login.tsx": "React TypeScript source for the login page UI.",
    "login_ui_mockup.png": "Login screen screenshot/mockup used as UI evidence.",
    "AuthContext.tsx": "React context source implementing login, session persistence, and role information.",
    "Books.tsx": "React TypeScript source for book catalog/search and book-management screens.",
    "book_catalog_1785261304299.png": "Book catalog screenshot showing catalog search and availability UI.",
    "Project_Planning_and_Structure.docx": "Planning document describing project structure and management approach.",
    "Requirements_Analysis_SRS.docx": "Formal requirements/SRS document uploaded as requirements evidence.",
    "SRS.md": "Markdown SRS containing user roles, functional requirements, and acceptance criteria.",
    "use_case_diagram.png": "StarUML-style use-case diagram showing Admin, Librarian, and Student interactions.",
    "WBS.md": "Work Breakdown Structure document mapping LMS work into project phases.",
    "library_management_system_codebase.zip": "Compressed project codebase uploaded as project-structure evidence.",
    "WhatsApp Image 2026-07-27 at 9.49.54 PM.jpeg": "Handwritten project-planning photo showing manual planning notes.",
    "WhatsApp Image 2026-07-27 at 9.49.56 PM.jpeg": "Handwritten project-planning photo showing manual planning notes.",
    "WhatsApp Image 2026-07-27 at 9.49.57 PM.jpeg": "Handwritten project-planning photo showing manual planning notes.",
    "WhatsApp Image 2026-07-27 at 9.50.00 PM.jpeg": "Handwritten project-planning photo showing manual planning notes.",
    "MyLoans.tsx": "React TypeScript source for the student loans/reservations view.",
    "checkout_activity_diagram.png": "StarUML-style activity diagram showing the book checkout workflow.",
    "library_management_system_advanced_features.zip": "Compressed bundle of advanced LMS feature files.",
    "LibraryContext.tsx": "React context source exposing book, member, issue/return, fine, settings, and reservation APIs.",
    "return_activity_diagram.png": "StarUML-style activity diagram showing the book return workflow.",
    "storage.ts": "LocalStorage database service containing mock relational data and transaction logic.",
    "Dashboard.tsx": "React TypeScript source for dashboard metrics and operational summaries.",
    "Project_Status_Snapshot.md": "Status snapshot documenting implementation progress and Jira synchronization.",
    "Settings.tsx": "React TypeScript source for administrative borrowing and fine policy settings.",
}


def parse_dt(value: str | None) -> str:
    if not value:
        return "Not recorded"
    for fmt in ("%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).strftime("%d %b %Y")
        except ValueError:
            continue
    return value[:10]


def completion(issue: dict[str, Any]) -> tuple[str, str]:
    if issue["status"] == "Done":
        return "100%", "Done in Jira"
    if issue.get("doneTransitions"):
        who = issue["doneTransitions"][-1]["by"]
        when = parse_dt(issue["doneTransitions"][-1]["at"])
        return "Work evidenced; currently reopened", f"Earlier Done transition by {who} on {when}; current status is {issue['status']}"
    return "In progress", f"Not marked Done in Jira; current status is {issue['status']}"


def finished_by(issue: dict[str, Any]) -> str:
    if issue.get("doneTransitions"):
        dt = issue["doneTransitions"][-1]
        return f"{dt['by']} on {parse_dt(dt['at'])}"
    if issue["status"] == "Done":
        return "Done in Jira; finisher not recorded in changelog"
    return "Not recorded as finished"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: int = 9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)


def style_table(table, header_fill="F2F4F7"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = RGBColor(46, 116, 181) if level <= 2 else RGBColor(31, 77, 120)
    return p


def add_body(doc, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_picture_if_exists(doc, path: str | Path, caption: str, width: float = 5.6):
    path = Path(path)
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(85, 85, 85)


def attachment_desc(filename: str) -> str:
    return ATTACHMENT_DESCRIPTIONS.get(filename, "Jira attachment uploaded as supporting evidence for this issue.")


def issue_sort_key(issue: dict[str, Any]) -> int:
    return int(issue["key"].split("-")[1])


def add_issue_section(doc: Document, issue: dict[str, Any]):
    add_heading(doc, f"{issue['key']} - {issue['summary']}", 2)
    pct, note = completion(issue)
    table = doc.add_table(rows=5, cols=4)
    table.style = "Table Grid"
    widths = [1.1, 2.0, 1.35, 2.05]
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
    pairs = [
        ("Type", issue["type"], "Status", issue["status"]),
        ("Assignee", issue.get("assignee") or "Unassigned", "Reporter", issue.get("reporter") or "Not recorded"),
        ("Priority", issue.get("priority") or "Not set", "Due date", parse_dt(issue.get("dueDate"))),
        ("Finished by", finished_by(issue), "Completion", pct),
        ("Attachment count", str(len(issue.get("attachments") or [])), "Completion note", note),
    ]
    for row, vals in zip(table.rows, pairs):
        for i, val in enumerate(vals):
            set_cell_text(row.cells[i], val, bold=i % 2 == 0, size=8)
            if i % 2 == 0:
                set_cell_shading(row.cells[i], "F2F4F7")
    style_table(table)

    add_body(doc, PROJECT_DESCRIPTIONS.get(issue["key"], "Jira work item recorded for the LMS project."))

    if issue.get("attachments"):
        add_heading(doc, "Attachments Recorded in Jira", 3)
        att_table = doc.add_table(rows=1, cols=3)
        att_table.style = "Table Grid"
        for i, h in enumerate(["Attachment", "Uploaded by/date", "Brief description"]):
            set_cell_text(att_table.rows[0].cells[i], h, bold=True, size=8)
        for att in issue["attachments"]:
            row = att_table.add_row().cells
            set_cell_text(row[0], att["filename"], size=8)
            set_cell_text(row[1], f"{att.get('author') or 'Not recorded'}; {parse_dt(att.get('created'))}", size=8)
            set_cell_text(row[2], attachment_desc(att["filename"]), size=8)
        style_table(att_table)
    else:
        add_body(doc, "Attachments recorded in Jira: none.")

    # Embed only the user-requested visual evidence: four handwritten planning photos and the StarUML-style diagrams.
    visual_names = {
        "WhatsApp Image 2026-07-27 at 9.49.54 PM.jpeg",
        "WhatsApp Image 2026-07-27 at 9.49.56 PM.jpeg",
        "WhatsApp Image 2026-07-27 at 9.49.57 PM.jpeg",
        "WhatsApp Image 2026-07-27 at 9.50.00 PM.jpeg",
        "use_case_diagram.png",
        "er_diagram.png",
        "checkout_activity_diagram.png",
        "return_activity_diagram.png",
    }
    visuals = [a for a in issue.get("attachments", []) if a["filename"] in visual_names]
    if visuals:
        add_heading(doc, "Embedded Visual Evidence", 3)
        for att in visuals:
            add_picture_if_exists(doc, att["localPath"], f"{issue['key']} attachment: {att['filename']}", width=5.7)


def main():
    data = json.loads(SNAPSHOT.read_text(encoding="utf-8-sig"))
    issues = sorted(data["issues"], key=issue_sort_key)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 12, RGBColor(31, 77, 120)),
    ]:
        st = styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(5)

    # Page 1 deliberately blank.
    doc.add_paragraph("")
    doc.add_page_break()

    # Page 2 starts project explanation, per user instruction.
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Library Management System\nJira Project Management Report")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(11, 37, 69)
    sub = doc.add_paragraph("Jira-focused project evidence report")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(11)
    sub.runs[0].font.color.rgb = RGBColor(85, 85, 85)

    add_heading(doc, "Project Overview", 1)
    add_body(
        doc,
        "The project managed in Jira is a Library Management System (LMS). The system objective is to support core library operations such as maintaining a book catalogue, allowing user authentication, managing members, issuing and returning books, handling fines, and showing dashboard/report information for the library workflow.",
    )
    add_body(
        doc,
        "This report is based on the Jira project named Library Management System Team with project key SCRUM. It focuses on the Jira sprint, epics, tasks, assignees, completion status, and evidence attachments recorded in Jira.",
    )

    add_heading(doc, "Jira Project Snapshot", 1)
    total = len(issues)
    done = sum(1 for i in issues if i["status"] == "Done")
    in_prog = sum(1 for i in issues if i["status"] == "In Progress")
    epics = sum(1 for i in issues if i["type"] == "Epic")
    tasks = sum(1 for i in issues if i["type"] == "Task")
    subtasks = sum(1 for i in issues if i["type"] == "Subtask")
    add_bullet(doc, f"Jira site: {data['project']['url']}")
    add_bullet(doc, f"Project key/name: {data['project']['key']} - {data['project']['name']}")
    add_bullet(doc, f"Board: SCRUM board")
    add_bullet(doc, f"Total Jira issues reviewed: {total} ({epics} epics, {tasks} tasks, {subtasks} subtasks)")
    add_bullet(doc, f"Current status count: {done} Done, {in_prog} In Progress")

    add_heading(doc, "Sprint Explanation", 1)
    sprint = None
    for board in data.get("sprintsByBoard", []):
        for sp in board.get("sprints", []):
            if sp.get("state") == "active":
                sprint = sp
                break
    if sprint:
        add_body(
            doc,
            f"The project work was organized around one active Jira sprint: {sprint['name']}. The sprint started on {parse_dt(sprint.get('startDate'))} and is scheduled to end on {parse_dt(sprint.get('endDate'))}. It belongs to the SCRUM board and is used to track the active development and documentation work for the LMS.",
        )
    else:
        add_body(doc, "The Jira board did not expose an active sprint in the API snapshot; issue work is therefore reported directly from the SCRUM board.")
    sprint_issues = [i for i in issues if i.get("sprint")]
    add_body(doc, f"Number of issues directly listed under the active sprint field: {len(sprint_issues)}.")

    summary_table = doc.add_table(rows=1, cols=5)
    summary_table.style = "Table Grid"
    for i, h in enumerate(["Key", "Type", "Summary", "Assignee", "Completion"]):
        set_cell_text(summary_table.rows[0].cells[i], h, bold=True, size=8)
    for issue in issues:
        pct, _ = completion(issue)
        row = summary_table.add_row().cells
        set_cell_text(row[0], issue["key"], size=8)
        set_cell_text(row[1], issue["type"], size=8)
        set_cell_text(row[2], issue["summary"], size=8)
        set_cell_text(row[3], issue.get("assignee") or "Unassigned", size=8)
        set_cell_text(row[4], pct, size=8)
    style_table(summary_table)

    doc.add_page_break()
    add_heading(doc, "Epics, Tasks, Assignees, Completion, and Attachments", 1)
    add_body(
        doc,
        "The following section records each Jira item exactly as a project-management artifact: issue key, issue type, assignee, current status, finishing evidence from Jira status history, completion state, and the attachments uploaded to that issue.",
    )

    for issue in issues:
        add_issue_section(doc, issue)

    add_heading(doc, "Overall Completion Based on Jira", 1)
    add_body(
        doc,
        f"Out of {total} Jira issues, {done} are currently marked Done and {in_prog} remain In Progress. Several issues have earlier Done transitions in the changelog but were later moved back to In Progress; these are reported as work evidenced but currently reopened. No extra completion claim has been added beyond the Jira status and history records.",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
